from odoo import fields,models,api
from odoo.exceptions import UserError
import openpyxl
import base64
from io import BytesIO


class UpdateProductSerial(models.TransientModel):
    _name = 'update.product.serial'
    _description = 'Update Product Serial'

    operation_type = fields.Selection([('replace_product','Replace Product'),('update_serial','Update Serial')])
    upload_file = fields.Binary(string='Upload File')
    replace_product_read_file_ids = fields.One2many('replace.product.read.file','update_product_serial_id', string='Replace Product File')
    replace_serial_read_file_ids = fields.One2many('replace.serial.read.file','update_product_serial_id', string='Replace Serial File')

    def read_product_file(self):
        wb = openpyxl.load_workbook(
            filename=BytesIO(base64.b64decode(self.upload_file)), read_only=True
        )
        ws = wb.active
        # print(ws)
        data = []
        self.replace_product_read_file_ids.unlink()

        for record in ws.iter_rows(min_row=2, max_row=None, min_col=None,
                                   max_col=None, values_only=True):
            mo_name = record[0].rstrip()
            current_product = self.env['product.product'].search([('default_code', '=', record[1])], limit=1)
            new_product = self.env['product.product'].search([('default_code', '=', record[2])], limit=1)
            mo_order = self.env['mrp.production'].search([('display_name','=', mo_name)], limit=1)
            data.append((0,0,{
                'mo_number_id': mo_order.id,
                'current_product':current_product.id,
                'new_product':new_product.id,
                'state':mo_order.state
            }))
        self.replace_product_read_file_ids = data
        view_id = self.env.ref('bitsa_manufacture.update_product_serial_form_view').id
        return {
            'name': 'Update Product Serial',
            'view_mode': 'form',
            'res_model': 'update.product.serial',
            'res_id': self.id,
            'view_id': view_id,
            'type': 'ir.actions.act_window',
            'target': 'new',
        }

    def replace_products(self):
        wb = openpyxl.load_workbook(
            filename=BytesIO(base64.b64decode(self.upload_file)), read_only=True
        )
        ws = wb.active
        for record in ws.iter_rows(min_row=2, max_row=None, min_col=None,
                                   max_col=None, values_only=True):
            mo_name = record[0].rstrip()
            mo_order = self.env['mrp.production'].search([('display_name', '=', mo_name)], limit=1)
            current_product = self.env['product.product'].search([('default_code', '=', record[1])], limit=1)
            new_product = self.env['product.product'].search([('default_code', '=', record[2])], limit=1)

            if mo_order and mo_order.product_id == current_product:
                mo_order.product_id = new_product

            elif mo_order and mo_order.product_id == new_product:
                mo_order.product_id = current_product

    def read_serial_file(self):
        wb = openpyxl.load_workbook(
            filename=BytesIO(base64.b64decode(self.upload_file)), read_only=True
        )
        ws = wb.active
        # print(ws)
        data = []
        self.replace_serial_read_file_ids.unlink()

        for record in ws.iter_rows(min_row=2, max_row=None, min_col=None,
                                   max_col=None, values_only=True):
            mo_name = record[0].rstrip()
            product_code = self.env['product.product'].search([('default_code', '=', record[1])], limit=1)
            old_serial = self.env['stock.lot'].search([('name','=',record[2])], limit=1)
            new_serial = self.env['stock.lot'].search([('name','=',record[3])], limit=1)
            mo_order = self.env['mrp.production'].search([('display_name', '=', mo_name)], limit=1)
            if not new_serial:
                # sequence = self.product_code.sequence_id.next_by_id()
                new_serial = self.env['stock.lot'].create({
                    'name':record[3],
                    'product_id': product_code.id,
                })
            data.append((0, 0, {
                'mo_number_id': mo_order.id,
                'product_code':product_code.id,
                'old_serial': old_serial.id,
                'new_serial': new_serial.id,
            }))
        self.replace_serial_read_file_ids = data
        view_id = self.env.ref('bitsa_manufacture.update_product_serial_form_view').id
        return {
            'name': 'Update Product Serial',
            'view_mode': 'form',
            'res_model': 'update.product.serial',
            'res_id': self.id,
            'view_id': view_id,
            'type': 'ir.actions.act_window',
            'target': 'new',
        }


    def replace_serial(self):
        wb = openpyxl.load_workbook(
            filename=BytesIO(base64.b64decode(self.upload_file)), read_only=True
        )
        ws = wb.active
        # print(ws)
        data = []

        for record in ws.iter_rows(min_row=2, max_row=None, min_col=None,
                                   max_col=None, values_only=True):
            mo_name = record[0].rstrip()
            product_code = self.env['product.product'].search([('default_code', '=', record[1])], limit=1)
            old_serial = self.env['stock.lot'].search([('name', '=', record[2])], limit=1)
            new_serial = self.env['stock.lot'].search([('name', '=', record[3])], limit=1)
            mo_order = self.env['mrp.production'].search([('display_name', '=', mo_name)], limit=1)
            if not new_serial:
                # sequence = self.product_code.sequence_id.next_by_id()
                new_serial = self.env['stock.lot'].create({
                    'name': record[3],
                    'product_id': product_code.id,
                })

            lot_ids = self.env['stock.move.line'].search([('lot_id', '=', old_serial.id)])

            if mo_order and lot_ids:
                lot_ids = new_serial






